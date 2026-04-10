"""
refresh_model_dropdown.py — 从 tao_config.xlsx 读取真实 API Key，
查询可用模型列表，刷新 Excel 中 MODEL_VIDEO / MODEL_IMAGE 的下拉选项。
"""
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

# 1. 读取 API Key
wb = openpyxl.load_workbook("tao_config.xlsx")
ws = wb["系统配置"]

cfg = {}
for row in ws.iter_rows(min_row=1, values_only=True):
    if row[0] and str(row[0]).strip() and not str(row[0]).startswith("←") and not str(row[0]).startswith("【"):
        cfg[str(row[0]).strip()] = str(row[1]).strip() if row[1] else ""

api_key = cfg.get("GCP_API_KEY", "").strip()
if not api_key:
    print("❌ tao_config.xlsx 中 GCP_API_KEY 为空，请先填写后再运行")
    exit(1)

print(f"API Key: {api_key[:12]}...")

# 2. 查询模型列表
from google import genai
client = genai.Client(api_key=api_key)

gemini_models = []
for m in client.models.list():
    name = m.name  # e.g. "models/gemini-2.5-pro"
    short = name.split("/")[-1]
    if "gemini" in short.lower():
        gemini_models.append(short)

gemini_models.sort()
print(f"找到 {len(gemini_models)} 个 Gemini 模型:")
for m in gemini_models:
    print(f"  {m}")

if not gemini_models:
    print("❌ 未找到任何 Gemini 模型，退出")
    exit(1)

# 3. 找到 MODEL_VIDEO / MODEL_IMAGE 所在行，更新单元格值
target_keys = {"MODEL_VIDEO", "MODEL_IMAGE"}
target_cells = []
for row in ws.iter_rows(min_row=1):
    key_cell = row[0]
    if key_cell.value and str(key_cell.value).strip() in target_keys:
        val_cell = row[1]
        target_cells.append(val_cell.coordinate)
        cur = str(val_cell.value).strip() if val_cell.value else ""
        if cur not in gemini_models:
            val_cell.value = gemini_models[0]
            print(f"  {key_cell.value}: reset '{cur}' -> '{gemini_models[0]}'")

# 4. 重建 DataValidation：保留非目标行的 DV，替换目标行的 DV
from openpyxl.worksheet.datavalidation import DataValidationList
dropdown_str = ",".join(gemini_models)

kept = []
for dv in ws.data_validations.dataValidation:
    sqref_str = str(dv.sqref)
    if not any(c in sqref_str for c in target_cells):
        kept.append(dv)

ws.data_validations = DataValidationList()
for dv in kept:
    ws.data_validations.append(dv)

for coord in target_cells:
    dv = DataValidation(
        type="list",
        formula1=f'"{dropdown_str}"',
        allow_blank=True,
        showDropDown=False,
    )
    dv.sqref = coord
    ws.data_validations.append(dv)
    print(f"  updated dropdown: {coord}")

wb.save("tao_config.xlsx")
print("Done. tao_config.xlsx saved.")
