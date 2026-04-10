import openpyxl
from openpyxl.styles import Font

data_raw = [
    ("9245", "jago.fang", "方梓杰", "jago.fang", "方梓杰"),
    ("9329", "koko.huang", "黄洁", "jago.fang", "方梓杰"),
    ("10111", "brain.zhang", "张桂萍", "jago.fang", "方梓杰"),
    ("10261", "jayce.liu", "刘成东", "jago.fang", "方梓杰"),
    ("9243", "jaydon.chow", "邹东航", "jago.fang", "方梓杰"),
    ("10548", "mago.liu", "刘文彬", "jago.fang", "方梓杰"),
    ("", "dory.chen", "陈冬婷", "jago.fang", "方梓杰"),
    ("", "tanya.liang", "梁媛", "jago.fang", "方梓杰"),
    ("10796", "ben.luo", "罗家俊", "jago.fang", "方梓杰"),
    ("", "tina.ma", "马天悦", "jago.fang", "方梓杰"),
    ("", "roxy.ying", "英嘉维", "jago.fang", "方梓杰"),
    ("", "lila.yuan", "袁理馨", "jago.fang", "方梓杰"),
    ("10935", "theo.liu", "刘萄", "jago.fang", "方梓杰"),
    ("10936", "irene.chen", "陈烨茵", "jago.fang", "方梓杰"),
    ("9246", "yoyo.yuan", "袁媛", "ivan.ye", "叶浪"),
    ("10623", "clairo.li", "李娜", "ivan.ye", "叶浪"),
    ("9425", "leely", "邱洁", "ivan.ye", "叶浪"),
    ("", "valentina.liao", "廖天爱", "ivan.ye", "叶浪"),
    ("", "junjie.huang", "黄俊杰", "ivan.ye", "叶浪"),
    ("", "regina.huang", "黄思瑶", "ivan.ye", "叶浪"),
    ("", "hei.wu", "吴令希", "ivan.ye", "叶浪"),
    ("", "xu.xu", "徐增棚", "ivan.ye", "叶浪"),
    ("", "lorena.zhang", "张雅璐", "ivan.ye", "叶浪"),
    ("9465", "east.nie", "聂慧东", "ivan.ye", "叶浪"),
    ("9419", "mathilda.wu", "吴晓婷", "ivan.ye", "叶浪"),
    ("10969", "enki.hu", "胡颖之", "ivan.ye", "叶浪"),
    ("10991", "vinnona.sun", "孙子涵", "ivan.ye", "叶浪"),
    ("9244", "ivan.ye", "叶浪", "ivan.ye", "叶浪"),
]

active = [(pid, eng, cn, sup_eng, sup_cn) for pid, eng, cn, sup_eng, sup_cn in data_raw if pid.strip()]

wb = openpyxl.load_workbook("arkclaw_config.xlsx")
ws = wb["设计师表"]

# 更新表头，增加主管列
ws.cell(1, 5).value = "主管姓名"
ws.cell(1, 6).value = "主管英文名"

# 清空旧数据行
for row in ws.iter_rows(min_row=3, max_row=ws.max_row):
    for cell in row:
        cell.value = None

body_font = Font(name="微软雅黑", size=10)

for i, (pid, eng_name, cn_name, sup_eng, sup_cn) in enumerate(active, 3):
    ws.row_dimensions[i].height = 22
    for col, val in enumerate([cn_name, pid, eng_name, "TRUE", sup_cn, sup_eng], 1):
        c = ws.cell(i, col, val)
        c.font = body_font

wb.save("arkclaw_config.xlsx")
print("写入 %d 条（过滤 %d 条无PID离职人员）" % (len(active), len(data_raw) - len(active)))
for pid, eng, cn, sup_eng, sup_cn in active:
    print("  %s  %s  %s  主管:%s" % (pid.ljust(6), cn.ljust(8), eng.ljust(20), sup_cn))
